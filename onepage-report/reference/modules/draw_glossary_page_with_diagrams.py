"""術語頁面（有圖片版 6 格）"""
from pptx.util import Inches, Pt

from ._colors import COLOR_TEXT, FONT_NAME
from .draw_glossary_card_with_diagram import draw_glossary_card_with_diagram


def draw_glossary_page_with_diagrams(slide, title, terms):
    """
    繪製一頁 6 格有圖片的術語卡片（2 列 x 3 欄）

    Args:
        slide: 投影片物件
        title: 頁面標題
        terms: 最多 6 個術語，每個是 {
            "term": "術語名稱",
            "desc": "解釋",
            "diagram_type": "flow|before_after|layers|timeline|icon",
            "diagram_params": {...}
        }
    """
    title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), Inches(12.7), Inches(0.4))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT
    p.font.name = FONT_NAME

    margin = 0.3
    gap = 0.2
    cols = 3
    rows = 2
    card_width = (13.333 - margin * 2 - gap * (cols - 1)) / cols
    card_height = (7.5 - 0.6 - gap * (rows - 1)) / rows

    for i, term_data in enumerate(terms[:6]):
        row = i // cols
        col = i % cols
        x = margin + col * (card_width + gap)
        y = 0.6 + row * (card_height + gap)

        draw_glossary_card_with_diagram(
            slide, x, y, card_width, card_height,
            term_data.get("term", ""),
            term_data.get("desc", ""),
            term_data.get("diagram_type", "icon"),
            term_data.get("diagram_params", {})
        )
