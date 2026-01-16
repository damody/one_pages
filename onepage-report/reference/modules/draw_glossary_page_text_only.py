"""術語頁面（純文字版 16 格）"""
from pptx.util import Inches, Pt

from ._colors import COLOR_TEXT, FONT_NAME
from .draw_glossary_card_text_only import draw_glossary_card_text_only


def draw_glossary_page_text_only(slide, title, terms):
    """
    繪製一頁 16 格純文字術語卡片（4 列 x 4 欄）

    Args:
        slide: 投影片物件
        title: 頁面標題
        terms: 最多 16 個術語，每個是 {
            "term": "術語名稱",
            "desc": "簡短解釋（<=50字）"
        }
    """
    title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(12.7), Inches(0.35))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT
    p.font.name = FONT_NAME

    margin = 0.25
    gap = 0.1
    cols = 4
    rows = 4
    card_width = (13.333 - margin * 2 - gap * (cols - 1)) / cols
    card_height = (7.5 - 0.5 - gap * (rows - 1)) / rows

    for i, term_data in enumerate(terms[:16]):
        row = i // cols
        col = i % cols
        x = margin + col * (card_width + gap)
        y = 0.5 + row * (card_height + gap)

        draw_glossary_card_text_only(
            slide, x, y, card_width, card_height,
            term_data.get("term", ""),
            term_data.get("desc", "")
        )
